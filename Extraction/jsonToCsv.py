import json
import csv
from pathlib import Path
from collections import OrderedDict
from decimal import Decimal, InvalidOperation
import re
import openpyxl
from openpyxl.utils import get_column_letter
import shutil
import copy
from compact_schema import STATEMENT_SUFFIX_ALTERNATION

# ── Coordinate support ────────────────────────────────────────────────────────
_COORD_SUFFIX = "_coord"

def _is_coord_key(key: str) -> bool:
    return key.endswith(_COORD_SUFFIX)

def _parent_of_coord(coord_key: str) -> str:
    return coord_key[: -len(_COORD_SUFFIX)]

def _coord_key_of(col: str) -> str:
    return col + _COORD_SUFFIX

def _serialise_coord(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return json.dumps(value, separators=(",", ":"), ensure_ascii=False)
    return str(value)

# ─────────────────────────────────────────────────────────────────────────────

try:
    from total_check_engine import run_total_check
except ImportError:
    def run_total_check(data, stmt_type):
        return data

CHECK_STATUS_COL = "Total Check Status"
ROW_PAGE_NO_COL  = "Row Page No"

# ── Column normalisation ──────────────────────────────────────────────────────
# Maps every statement-specific "* Items" key to the canonical "Row Items".
# This ensures all 15 sheets share one consistent label column name.
COLUMN_RENAME_MAP: dict[str, str] = {
    "SOA Items":            "Row Items",
    "GOV_BS Items":         "Row Items",
    "GOV_IS Items":         "Row Items",
    "PROP_SNP Items":       "Row Items",
    "PROP_IS Items":        "Row Items",
    "PROP_CFS Items":       "Row Items",
    "DSR Items":            "Row Items",
    "TAX_BASE Items":       "Row Items",
    "Pension Items":        "Row Items",
    "OPEB Items":           "Row Items",
    "CAPITAL_ASSETS Items": "Row Items",   # safety alias
    "Finding":              "Row Items",   # FAQS sheet
    "Overview Item":        "Row Items",   # OVERVIEW sheet
}

# Columns that duplicate the Section column and must be suppressed entirely.
COLUMNS_TO_DROP: set[str] = {
    "Pension Plan Name",
    "OPEB Plan Name",
}

# Reverse map: canonical name → original JSON key (first registered wins).
_RENAME_REVERSE: dict[str, str] = {}
for _src, _dst in COLUMN_RENAME_MAP.items():
    if _dst not in _RENAME_REVERSE:
        _RENAME_REVERSE[_dst] = _src
# ─────────────────────────────────────────────────────────────────────────────

_SELF_CHECKED_STMT_TYPES: set[str] = {
    "_CAPITAL_ASSETS",
    "_TAX_BASE",
    "_OVERVIEW",
    "_PEN",
    "_OPEB",
    "_FAQS",
}

_STMT_JSON_RE = re.compile(
    rf"^(.+?)_({STATEMENT_SUFFIX_ALTERNATION})"
    r"(?:_Comp)?_p[\d\-]+$",
    re.IGNORECASE,
)


def infer_stmt_type(json_stem: str) -> str:
    m = _STMT_JSON_RE.match(json_stem)
    if m:
        return "_" + m.group(2).upper()
    return ""


METADATA_COLUMNS = {
    "Statement", "Issuer Name", "FYE", "Page No", "Currency reported",
    "Reported in Thousand", "Section",
}

NON_NUMERIC_ITEM_COLUMNS = {
    "COA Flag",
    "Row Items",
    "COA Datapoint",
}


def parse_amount(value) -> Decimal:
    if value is None:
        return Decimal("0")
    s = str(value).strip()
    if s in {"", "-", "\u2013", "\u2014"} or set(s) == {"#"}:
        return Decimal("0")
    s = s.replace("$", "").replace(",", "").replace(" ", "")
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1]
    if s.startswith("-"):
        negative, s = True, s[1:]
    if s in {"", "-", "\u2013", "\u2014"}:
        return Decimal("0")
    try:
        num = Decimal(s)
    except InvalidOperation:
        return Decimal("0")
    return -num if negative else num


def format_amount(value: Decimal) -> str:
    if value is None:
        return ""
    value = Decimal(value)
    return f"{int(value):,}" if value == value.to_integral_value() else f"{value:,.2f}"


def is_numeric_reporting_column(col: str, csv_rows: list[dict]) -> bool:
    if col in METADATA_COLUMNS or col in NON_NUMERIC_ITEM_COLUMNS:
        return False
    for row in csv_rows:
        v = str(row.get(col) or "").strip()
        if not v or v in {"-", "\u2013", "\u2014"}:
            continue
        cleaned = re.sub(r"[$,\s()\-]", "", v)
        if cleaned.replace(".", "", 1).isdigit():
            return True
    return False


def get_row_label(row: dict) -> str:
    for key in ["Row Items"]:
        v = row.get(key)
        if v:
            return str(v).strip()
    return ""


def normalize_value(value):
    if value is None:
        return ""
    if isinstance(value, dict):
        return _serialise_coord(value)
    if isinstance(value, list):
        return json.dumps(value, ensure_ascii=False)
    return value


def _apply_column_renames(normal_cols: list[str]) -> list[str]:
    """
    1. Drop columns in COLUMNS_TO_DROP (Pension/OPEB Plan Name — duplicates Section).
    2. Rename columns via COLUMN_RENAME_MAP (e.g. "Pension Items" → "Row Items").
    3. De-duplicate while preserving first-seen order.
    """
    seen:   set[str]  = set()
    result: list[str] = []
    for col in normal_cols:
        if col in COLUMNS_TO_DROP:
            continue
        canonical = COLUMN_RENAME_MAP.get(col, col)
        if canonical not in seen:
            seen.add(canonical)
            result.append(canonical)
    return result


def discover_dynamic_columns(sections: dict) -> tuple[list[str], list[str]]:
    """
    Walk all section rows and return two lists:
      1. normal_cols  — every key that is NOT a _coord key, in discovery order
      2. coord_cols   — every _coord key found, in discovery order
    """
    normal_cols: OrderedDict[str, bool] = OrderedDict()
    coord_cols:  OrderedDict[str, bool] = OrderedDict()

    for _, items in sections.items():
        if not isinstance(items, list):
            continue
        for item in items:
            if not isinstance(item, dict):
                continue
            for key in item.keys():
                if _is_coord_key(key):
                    coord_cols[key] = True
                else:
                    normal_cols[key] = True

    return list(normal_cols.keys()), list(coord_cols.keys())


def order_item_columns(
    normal_item_columns: list[str],
    coord_item_columns:  list[str],
) -> list[str]:
    """
    Build the final ordered column list:
      • Priority label columns ALWAYS first, in this fixed sequence:
          COA Flag  →  Row Items  →  COA Datapoint
        COA Datapoint is ALWAYS emitted as column J (after the 7 PRIORITY_COLS
        metadata columns in A–G + COA Flag in H + Row Items in I).
        It is injected even when absent from the JSON data so that every sheet
        has a consistent column J = "COA Datapoint".
      • Then each reporting column immediately followed by its _coord sibling.
      • Total Check Status is always last.
    """
    # Fixed label columns — order determines Excel column positions H, I, J.
    # COA Datapoint is ALWAYS included (injected if missing from data) so it
    # lands in column J on every sheet without exception.
    priority = [
        "COA Flag",       # col H
        "Row Items",      # col I
        "COA Datapoint",  # col J  ← always present, always here
    ]

    coord_set = set(coord_item_columns)

    filtered = [c for c in normal_item_columns if c != CHECK_STATUS_COL]

    # Always place all priority columns first — inject any that are absent from
    # the discovered data so column J is never skipped.
    ordered   = list(priority)                                    # always all 3
    remaining = [c for c in filtered if c not in set(priority)]  # reporting cols

    for col in remaining:
        ordered.append(col)
        coord_key = _coord_key_of(col)
        if coord_key in coord_set:
            ordered.append(coord_key)

    already_added = set(ordered)
    for ck in coord_item_columns:
        if ck not in already_added:
            ordered.append(ck)

    return ordered


def run_json_to_csv_pipeline(json_path: str, csv_path: str) -> bool:
    json_path = Path(json_path)
    csv_path  = Path(csv_path)

    if not json_path.exists():
        raise FileNotFoundError(f"JSON file not found: {json_path}")

    with open(json_path, "r", encoding="utf-8", errors="replace") as f:
        data = json.load(f)

    metadata = data.get("Metadata", {})
    sections = data.get("Sections", {})

    if not isinstance(sections, dict) or not sections:
        raise ValueError("Invalid or empty 'Sections' found in JSON")

    stmt_type = infer_stmt_type(json_path.stem)
    if not stmt_type:
        print(f" [WARN] Could not infer statement type from filename "
              f"'{json_path.name}' — skipping Total Check arithmetic and "
              f"column-shift repair; all CP rows will read as "
              f"Total Check Status = \"\".",
              flush=True)
        checked_sections = sections
    else:
        try:
            from column_shift_repair import repair_column_shifts
            repair_column_shifts(data, stmt_type)
        except ImportError:
            print(" [WARN] column_shift_repair.py not found — skipping repair.",
                  flush=True)
        except Exception as e:
            print(f" [WARN] column_shift_repair failed: {e}", flush=True)

        if data.get("_column_shift_repairs"):
            json_path.write_text(
                json.dumps(data, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            print(f"   [REPAIR] {len(data['_column_shift_repairs'])} "
                  f"column-shift fix(es) written back to {json_path.name}",
                  flush=True)
            sections = data.get("Sections", {})

        if stmt_type in _SELF_CHECKED_STMT_TYPES:
            print(f"   [CHECK] {stmt_type.lstrip('_')} - Total Check Status kept "
                  f"as extracted (no total_check_engine rules for this type)",
                  flush=True)
            checked_sections = sections
        else:
            checked_data     = run_total_check(copy.deepcopy(data), stmt_type)
            checked_sections = checked_data.get("Sections", {})

    statement             = metadata.get("Statement", "")
    issuer_name           = metadata.get("Issuer Name", "")
    fye                   = metadata.get("FYE", "")
    page_no               = metadata.get("Page No", "")
    currency_reported     = metadata.get("Currency reported", "")
    reported_in_thousand  = metadata.get("Reported in Thousand", "No")

    # ── Discover columns (coords separated from normals) ─────────────────
    normal_cols, coord_cols = discover_dynamic_columns(checked_sections)

    # Rename statement-specific "* Items" → "Row Items" and drop duplicate
    # plan-name columns (Pension/OPEB Plan Name == Section).
    normal_cols = _apply_column_renames(normal_cols)

    ordered_item_columns = order_item_columns(normal_cols, coord_cols)

    # ── Resolve the "Page No" name collision ─────────────────────────────
    output_item_columns = [
        ROW_PAGE_NO_COL if c == "Page No" else c
        for c in ordered_item_columns
    ]
    if ROW_PAGE_NO_COL in output_item_columns:
        print(f"   [COLS] per-row page column renamed "
              f"\"Page No\" -> \"{ROW_PAGE_NO_COL}\" to protect the metadata "
              f"Page No identifier", flush=True)

    base_fieldnames = (
        ["Statement", "Issuer Name", "FYE", "Page No", "Currency reported",
         "Reported in Thousand", "Section"]
        + output_item_columns
        + [CHECK_STATUS_COL]
    )

    # ── Build CSV rows ────────────────────────────────────────────────────
    csv_rows = []
    for section_name, items in checked_sections.items():
        if not isinstance(items, list):
            continue
        for item in items:
            if not isinstance(item, dict):
                continue
            row = {
                "Statement":            statement,
                "Issuer Name":          issuer_name,
                "FYE":                  fye,
                "Page No":              page_no,
                "Currency reported":    currency_reported,
                "Reported in Thousand": reported_in_thousand,
                "Section":              section_name,
            }
            for src_col, out_col in zip(ordered_item_columns, output_item_columns):
                if _is_coord_key(src_col):
                    raw = item.get(src_col)
                    row[out_col] = _serialise_coord(raw)
                else:
                    # src_col is already the CANONICAL name (after _apply_column_renames).
                    # Look up the value using both the canonical name and any original
                    # JSON key that maps to it, so we find the value regardless of
                    # which key name the LLM used in the JSON.
                    raw = item.get(src_col)
                    if raw is None:
                        # Try original key names that map to this canonical name
                        for orig_key, canonical in COLUMN_RENAME_MAP.items():
                            if canonical == src_col:
                                raw = item.get(orig_key)
                                if raw is not None:
                                    break
                    row[out_col] = normalize_value(raw)
            row[CHECK_STATUS_COL] = normalize_value(item.get(CHECK_STATUS_COL))
            csv_rows.append(row)

    if not csv_rows:
        raise RuntimeError("No rows generated for CSV")

    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=base_fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(csv_rows)

    ACCEPTABLE = {"", "PASS", "Skipped - no membership rule defined",
                  "Skipped - no rule defined"}
    all_pass = all(
        row.get(CHECK_STATUS_COL, "") in ACCEPTABLE
        for row in csv_rows
        if str(row.get("COA Flag", "")).strip().upper() == "CP"
    )

    return all_pass


# ─────────────────────────────────────────────────────────────────────────────
# SHARED REGEX / CONSTANTS FOR MERGE FUNCTIONS
# ─────────────────────────────────────────────────────────────────────────────

_STMT_CSV_RE = re.compile(
    rf"^(.+?)_({STATEMENT_SUFFIX_ALTERNATION})"
    r"(?:_Comp)?_p[\d\-]+\.csv$",
    re.IGNORECASE,
)

_ALL_CSV_RE = re.compile(r"_All\.csv$", re.IGNORECASE)

STMT_ORDER = {
    "OVERVIEW":        0,
    "SNP":             1,
    "SOA":             2,
    "GOV_BS":          3,
    "GOV_IS":          4,
    "PROP_SNP":        5,
    "PROP_IS":         6,
    "PROP_CFS":        7,
    "CAPITAL_ASSETS":  8,
    "DEBT":            9,
    "DSR":            10,
    "TAX_BASE":       11,
    "PEN":            12,
    "OPEB":           13,
    "FAQS":           14,
}

PRIORITY_COLS = [
    "Statement", "Issuer Name", "FYE", "Page No",
    "Currency reported", "Reported in Thousand", "Section",
]


# ─────────────────────────────────────────────────────────────────────────────
# MERGE ALL CSVs PER ISSUER INTO ONE _All.csv
# ─────────────────────────────────────────────────────────────────────────────

def merge_all_csvs(folder: str, output_folder: str | None = None) -> list[str]:
    folder = Path(folder)
    if output_folder is None:
        output_folder = folder
    else:
        output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    groups: dict[str, list[Path]] = {}

    for f in sorted(folder.rglob("*.csv")):
        if not f.is_file():
            continue
        if _ALL_CSV_RE.search(f.name):
            continue
        m = _STMT_CSV_RE.match(f.name)
        if not m:
            continue
        base_name = m.group(1)
        if base_name not in groups:
            groups[base_name] = []
        groups[base_name].append(f)

    if not groups:
        print(f"  No statement CSVs found in: {folder}")
        return []

    for base_name in groups:
        groups[base_name].sort(
            key=lambda p: STMT_ORDER.get(
                (_STMT_CSV_RE.match(p.name).group(2).upper()
                 if _STMT_CSV_RE.match(p.name) else ""),
                99,
            )
        )

    created = []

    for base_name, csv_files in sorted(groups.items()):
        print(f"\n  Merging: {base_name}")

        all_columns = OrderedDict()
        all_rows    = []

        for csv_file in csv_files:
            stmt_m    = _STMT_CSV_RE.match(csv_file.name)
            stmt_type = stmt_m.group(2).upper() if stmt_m else "?"

            with open(csv_file, "r", encoding="utf-8-sig", errors="replace") as fh:
                reader    = csv.DictReader(fh)
                fieldnames = list(reader.fieldnames or [])
                rows      = list(reader)

            for col in fieldnames:
                if col not in all_columns:
                    all_columns[col] = True

            all_rows.extend(rows)
            print(f"    + {csv_file.name}  ({len(rows)} rows)  [{stmt_type}]")

        if not all_rows:
            print(f"    ⚠ No rows — skipping")
            continue

        final_columns = []
        for col in PRIORITY_COLS:
            if col in all_columns:
                final_columns.append(col)
        for col in all_columns:
            if col not in final_columns:
                final_columns.append(col)

        out_path = output_folder / f"{base_name}_All.csv"
        with open(out_path, "w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=final_columns, extrasaction="ignore")
            writer.writeheader()
            for row in all_rows:
                writer.writerow({col: row.get(col, "") for col in final_columns})

        print(f"    → {out_path.name}  ({len(all_rows)} total rows)")
        created.append(str(out_path))

    print(f"\n  Done. {len(created)} merged file(s) created.")
    return created


# ─────────────────────────────────────────────────────────────────────────────
# MERGE CSVs INSIDE ONE DEAL FOLDER → EXCEL (sheet per statement type)
# THEN DELETE INDIVIDUAL CSVs
# ─────────────────────────────────────────────────────────────────────────────

def merge_deal_csvs_to_excel(deal_folder: str, deal_name: str) -> str:
    deal_folder = Path(deal_folder)
    output_path = deal_folder / f"{deal_name}_All_Statements.xlsx"

    csv_files_by_type: dict[str, list[Path]] = {}
    all_csv_files:     list[Path]            = []

    for f in sorted(deal_folder.glob("*.csv")):
        if not f.is_file():
            continue
        if _ALL_CSV_RE.search(f.name):
            continue
        if f.name.startswith("_"):
            continue
        m = _STMT_CSV_RE.match(f.name)
        if not m:
            continue
        stmt_type = m.group(2).upper()
        csv_files_by_type.setdefault(stmt_type, []).append(f)
        all_csv_files.append(f)

    if not csv_files_by_type:
        print(f"    No statement CSVs found in: {deal_folder}")
        return ""

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ordered_types = sorted(
        csv_files_by_type.keys(),
        key=lambda t: STMT_ORDER.get(t, 99),
    )

    grey_fill  = openpyxl.styles.PatternFill("solid", fgColor="D9D9D9")
    coord_font = openpyxl.styles.Font(color="808080", italic=True, size=8)

    for stmt_type in ordered_types:
        csv_files = sorted(csv_files_by_type[stmt_type])

        all_columns = OrderedDict()
        all_rows    = []

        for csv_file in csv_files:
            with open(csv_file, "r", encoding="utf-8-sig", errors="replace") as fh:
                reader     = csv.DictReader(fh)
                fieldnames = list(reader.fieldnames or [])
                rows       = list(reader)

            for col in fieldnames:
                if col not in all_columns:
                    all_columns[col] = True

            all_rows.extend(rows)

        if not all_rows:
            continue

        final_columns = []
        for col in PRIORITY_COLS:
            if col in all_columns:
                final_columns.append(col)
        for col in all_columns:
            if col not in final_columns:
                final_columns.append(col)

        ws = wb.create_sheet(title=stmt_type)

        for col_idx, col_name in enumerate(final_columns, start=1):
            cell      = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = openpyxl.styles.Font(bold=True)
            if _is_coord_key(col_name):
                cell.fill = grey_fill
                cell.font = openpyxl.styles.Font(bold=True, color="505050", size=8)

        for row_idx, row_data in enumerate(all_rows, start=2):
            for col_idx, col_name in enumerate(final_columns, start=1):
                value = row_data.get(col_name, "")
                cell  = ws.cell(row=row_idx, column=col_idx, value=value)
                if _is_coord_key(col_name) and value:
                    cell.font = coord_font
                    cell.fill = openpyxl.styles.PatternFill("solid", fgColor="F5F5F5")

        for col_idx, col_name in enumerate(final_columns, start=1):
            max_len = len(col_name)
            for row_data in all_rows[:50]:
                cell_val = str(row_data.get(col_name, ""))
                max_len  = max(max_len, len(cell_val))
            if _is_coord_key(col_name):
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 35)
                ws.column_dimensions[get_column_letter(col_idx)].hidden = True
            else:
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    if not wb.sheetnames:
        print(f"    No data sheets created — skipping Excel save for {deal_name}")
        return ""

    wb.save(str(output_path))
    print(f"    ✅ {output_path.name}")

    deleted = 0
    for csv_file in all_csv_files:
        try:
            csv_file.unlink()
            deleted += 1
        except Exception as e:
            print(f"    [WARN] Could not delete {csv_file.name}: {e}")

    return str(output_path)


# ─────────────────────────────────────────────────────────────────────────────
# MERGE ALL CSVs INTO ONE EXCEL — SHEET PER STATEMENT TYPE
# ─────────────────────────────────────────────────────────────────────────────

def merge_csvs_to_excel(folder: str, output_path: str | None = None) -> str:
    folder = Path(folder)

    if output_path is None:
        output_path = folder / "All_Statements.xlsx"
    else:
        output_path = Path(output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)

    stmt_groups: dict[str, list[Path]] = {}

    for f in sorted(folder.rglob("*.csv")):
        if not f.is_file():
            continue
        if _ALL_CSV_RE.search(f.name):
            continue
        if f.name.startswith("_"):
            continue
        m = _STMT_CSV_RE.match(f.name)
        if not m:
            continue
        stmt_type = m.group(2).upper()
        if stmt_type not in stmt_groups:
            stmt_groups[stmt_type] = []
        stmt_groups[stmt_type].append(f)

    if not stmt_groups:
        print(f"  No statement CSVs found in: {folder}")
        return ""

    wb           = openpyxl.Workbook()
    grey_fill    = openpyxl.styles.PatternFill("solid", fgColor="D9D9D9")
    coord_font   = openpyxl.styles.Font(color="808080", italic=True, size=8)
    wb.remove(wb.active)

    ordered_types = sorted(stmt_groups.keys(), key=lambda t: STMT_ORDER.get(t, 99))

    for stmt_type in ordered_types:
        csv_files = sorted(stmt_groups[stmt_type])
        print(f"\n  Sheet: {stmt_type}  ({len(csv_files)} CSV file(s))")

        all_columns = OrderedDict()
        all_rows    = []

        for csv_file in csv_files:
            with open(csv_file, "r", encoding="utf-8-sig", errors="replace") as fh:
                reader     = csv.DictReader(fh)
                fieldnames = list(reader.fieldnames or [])
                rows       = list(reader)

            for col in fieldnames:
                if col not in all_columns:
                    all_columns[col] = True

            all_rows.extend(rows)
            print(f"    + {csv_file.name}  ({len(rows)} rows)")

        if not all_rows:
            print(f"    ⚠ No rows — skipping sheet")
            continue

        final_columns = []
        for col in PRIORITY_COLS:
            if col in all_columns:
                final_columns.append(col)
        for col in all_columns:
            if col not in final_columns:
                final_columns.append(col)

        ws = wb.create_sheet(title=stmt_type)

        for col_idx, col_name in enumerate(final_columns, start=1):
            cell      = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = openpyxl.styles.Font(bold=True)
            if _is_coord_key(col_name):
                cell.fill = grey_fill
                cell.font = openpyxl.styles.Font(bold=True, color="505050", size=8)

        for row_idx, row_data in enumerate(all_rows, start=2):
            for col_idx, col_name in enumerate(final_columns, start=1):
                value = row_data.get(col_name, "")
                cell  = ws.cell(row=row_idx, column=col_idx, value=value)
                if _is_coord_key(col_name) and value:
                    cell.font = coord_font
                    cell.fill = openpyxl.styles.PatternFill("solid", fgColor="F5F5F5")

        for col_idx, col_name in enumerate(final_columns, start=1):
            max_len = len(col_name)
            for row_data in all_rows[:50]:
                cell_val = str(row_data.get(col_name, ""))
                max_len  = max(max_len, len(cell_val))
            if _is_coord_key(col_name):
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 35)
            else:
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

        print(f"    → Sheet '{stmt_type}': {len(all_rows)} rows, {len(final_columns)} columns")

    if not wb.sheetnames:
        print(f"  No data sheets created — skipping Excel save")
        return ""

    wb.save(str(output_path))
    print(f"\n  ✅ Excel saved: {output_path}")
    return str(output_path)


# ─────────────────────────────────────────────────────────────────────────────
# CLI / STANDALONE
# ─────────────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    DEAL_FOLDER = r"C:\Users\sbusr1\Downloads\Json\Json"

    MANUAL_VALIDATION_ROOT = (
        Path(DEAL_FOLDER).parent.parent / "05_Manual_Validation_Required"
    )

    deal_path = Path(DEAL_FOLDER)
    deal_name = deal_path.name

    if not deal_path.exists():
        print(f"[ERROR] Folder not found: {deal_path}")
        exit(1)

    json_files = sorted(deal_path.glob("*.json"))

    if not json_files:
        print(f"[WARN] No JSON files found in: {deal_path}")
        exit(1)

    print(f"{'='*70}")
    print(f"  Deal   : {deal_name}")
    print(f"  Folder : {deal_path}")
    print(f"  JSONs  : {len(json_files)}")
    print(f"{'='*70}\n")

    success, failed = 0, 0
    failed_tables: list[str] = []

    for jf in json_files:
        csv_path = deal_path / (jf.stem + ".csv")
        print(f"  [{jf.name}]")
        try:
            result = run_json_to_csv_pipeline(str(jf), str(csv_path))
            status = "PASS" if result else "FAIL(s) detected"
            print(f"    -> {csv_path.name}  ({status})\n")
            success += 1
            if not result:
                failed_tables.append(jf.stem)
        except Exception as e:
            print(f"    [ERROR] {e}\n")
            failed += 1
            failed_tables.append(jf.stem)

    print(f"\n  JSON → CSV : {success} succeeded, {failed} failed")
    if failed_tables:
        print(f"  Validation FAIL on : {failed_tables}\n")
    else:
        print(f"  Validation : ALL PASS\n")

    print(f"  Merging CSVs into Excel...")
    xlsx_path = merge_deal_csvs_to_excel(str(deal_path), deal_name)

    if xlsx_path:
        print(f"\n  ✅ Excel : {xlsx_path}")
    else:
        print(f"\n  [WARN] No Excel created")

    if failed_tables:
        MANUAL_VALIDATION_ROOT.mkdir(parents=True, exist_ok=True)
        target_dir = MANUAL_VALIDATION_ROOT / deal_name

        if target_dir.exists():
            shutil.rmtree(target_dir)

        try:
            shutil.move(str(deal_path), str(target_dir))
            print(f"\n  ⚠  FAIL detected in: {failed_tables}")
            print(f"  📦 Deal moved to   : {target_dir}")
        except Exception as e:
            print(f"\n  [ERROR] Could not move deal to manual validation: {e}")
    else:
        print(f"\n  ✅ All tables passed — deal stays in 04_Validated_Output")

    print(f"\n{'='*70}\n  DONE\n{'='*70}")